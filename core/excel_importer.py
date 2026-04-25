"""
core/excel_importer.py
──────────────────────
Xử lý nhập/xuất file Excel đáp án trắc nghiệm.
  • import_excel()          : đọc file Excel đáp án → dict {ma_de: {q: ans}}
  • create_template_bytes() : tạo file Excel mẫu dưới dạng bytes
"""

import io

VALID_ANSWERS = {"A", "B", "C", "D"}
NUM_QUESTIONS  = 40


def import_excel(path: str) -> tuple[dict, list[str]]:
    """
    Đọc file Excel đáp án trắc nghiệm.

    Returns
    -------
    result : dict
        {ma_de: {q_num(int): ans(str)}}
    errors : list[str]
        Danh sách cảnh báo / lỗi từng dòng
    """
    import pandas as pd

    errors: list[str] = []
    result: dict = {}

    try:
        df = pd.read_excel(path, dtype=str)
    except Exception as e:
        raise ValueError(f"Không thể đọc file Excel: {e}")

    df.columns = [str(c).strip() for c in df.columns]
    made_col = next(
        (c for c in df.columns if c.lower() in ("ma_de", "made", "mã đề", "ma de")), None)
    if made_col is None:
        raise ValueError("Không tìm thấy cột mã đề.\nCột đầu tiên phải có tên: ma_de")

    for row_idx, row in df.iterrows():
        made = str(row[made_col]).strip().zfill(3)
        if not made or made == "nan":
            errors.append(f"Hàng {row_idx+2}: bỏ qua vì mã đề trống")
            continue
        key_dict: dict[int, str] = {}
        for q in range(1, NUM_QUESTIONS + 1):
            col_name = str(q)
            if col_name not in df.columns:
                errors.append(f"Hàng {row_idx+2} mã '{made}': thiếu cột câu {q}, mặc định 'A'")
                key_dict[q] = "A"
                continue
            val = str(row[col_name]).strip().upper()
            key_dict[q] = val if val in VALID_ANSWERS else "A"
            if val not in VALID_ANSWERS:
                errors.append(f"Hàng {row_idx+2} mã '{made}' câu {q}: '{val}' không hợp lệ → 'A'")
        result[made] = key_dict

    return result, errors


def create_template_bytes() -> bytes:
    """
    Tạo file Excel mẫu chứa header + 3 dòng ví dụ.

    Returns
    -------
    bytes : nội dung file .xlsx dưới dạng bytes
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Đáp Án TN"

    hf = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    hc = PatternFill("solid", start_color="2980B9")
    ca = Alignment(horizontal="center", vertical="center")

    ws.cell(row=1, column=1, value="ma_de").font = hf
    ws.cell(row=1, column=1).fill = hc
    ws.cell(row=1, column=1).alignment = ca
    ws.column_dimensions["A"].width = 10

    for q in range(1, NUM_QUESTIONS + 1):
        c = ws.cell(row=1, column=q+1, value=q)
        c.font = hf; c.fill = hc; c.alignment = ca
        ws.column_dimensions[c.column_letter].width = 5

    sample = ["A", "B", "C", "D"]
    for r, made in enumerate(["001", "002", "003"], start=2):
        ws.cell(row=r, column=1, value=made).alignment = ca
        for q in range(1, NUM_QUESTIONS+1):
            ws.cell(row=r, column=q+1, value=sample[(q-1)%4]).alignment = ca

    ws.freeze_panes = "B2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
