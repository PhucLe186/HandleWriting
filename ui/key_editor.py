"""
app/ui/key_editor.py
────────────────────
Cửa sổ popup quản lý mã đề & đáp án — 2 tab:
  • Tab Trắc nghiệm : 40 câu A/B/C/D + import Excel
  • Tab Tự luận     : N câu, mỗi câu có đáp án mẫu text + điểm tối đa
"""

import io
import os
from tkinter import filedialog, messagebox
from typing import Callable

import customtkinter as ctk

VALID_ANSWERS = {"A", "B", "C", "D"}
NUM_QUESTIONS  = 40
NUM_TL_DEFAULT = 5      # số câu tự luận mặc định

# ── Excel helpers (trắc nghiệm) ───────────────────────────────────────────────

def _import_excel(path: str) -> tuple[dict, list[str]]:
    import pandas as pd
    errors: list[str] = []
    result: dict = {}
    try:
        df = pd.read_excel(path, dtype=str)
    except Exception as e:
        raise ValueError(f"Không thể đọc file Excel: {e}")
    df.columns = [str(c).strip() for c in df.columns]
    made_col = next(
        (c for c in df.columns if c.lower() in ("ma_de", "made", "mã đề", "ma de")), None)
    if made_col is None:
        raise ValueError("Không tìm thấy cột mã đề.\nCột đầu tiên phải có tên: ma_de")
    for row_idx, row in df.iterrows():
        made = str(row[made_col]).strip().zfill(3)
        if not made or made == "nan":
            errors.append(f"Hàng {row_idx+2}: bỏ qua vì mã đề trống")
            continue
        key_dict: dict[int, str] = {}
        for q in range(1, NUM_QUESTIONS + 1):
            col_name = str(q)
            if col_name not in df.columns:
                errors.append(f"Hàng {row_idx+2} mã '{made}': thiếu cột câu {q}, mặc định 'A'")
                key_dict[q] = "A"
                continue
            val = str(row[col_name]).strip().upper()
            key_dict[q] = val if val in VALID_ANSWERS else "A"
            if val not in VALID_ANSWERS:
                errors.append(f"Hàng {row_idx+2} mã '{made}' câu {q}: '{val}' không hợp lệ → 'A'")
        result[made] = key_dict
    return result, errors


def _create_template_bytes() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Đáp Án TN"
    hf = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    hc = PatternFill("solid", start_color="2980B9")
    ca = Alignment(horizontal="center", vertical="center")
    ws.cell(row=1, column=1, value="ma_de").font = hf
    ws.cell(row=1, column=1).fill = hc
    ws.cell(row=1, column=1).alignment = ca
    ws.column_dimensions["A"].width = 10
    for q in range(1, NUM_QUESTIONS + 1):
        c = ws.cell(row=1, column=q+1, value=q)
        c.font = hf; c.fill = hc; c.alignment = ca
        ws.column_dimensions[c.column_letter].width = 5
    sample = ["A","B","C","D"]
    for r, made in enumerate(["001","002","003"], start=2):
        ws.cell(row=r, column=1, value=made).alignment = ca
        for q in range(1, NUM_QUESTIONS+1):
            ws.cell(row=r, column=q+1, value=sample[(q-1)%4]).alignment = ca
    ws.freeze_panes = "B2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Main window ───────────────────────────────────────────────────────────────

class KeyEditorWindow(ctk.CTkToplevel):
    """
    Parameters
    ----------
    master         : cửa sổ cha
    answer_key_db  : {ma_de: {q_num(int): ans(str)}}
    tl_key_db      : {ma_de: [ {q: int, dap_an: str, diem: float} ]}
    on_save        : callback(ma_de, tn_dict, tl_list)
    on_bulk_save   : callback(tn_db) — sau import Excel
    """

    def __init__(
        self,
        master,
        answer_key_db: dict,
        tl_key_db: dict,
        on_save: Callable[[str, dict, list], None],
        on_bulk_save: Callable[[dict], None] | None = None,
    ):
        super().__init__(master)
        self.answer_key_db = answer_key_db
        self.tl_key_db     = tl_key_db
        self.on_save       = on_save
        self.on_bulk_save  = on_bulk_save

        self.title("Quản Lý Mã Đề & Đáp Án")
        self.geometry("1200x820")
        self.transient(master)
        self.update_idletasks()
        x = (self.winfo_screenwidth()  - 1200) // 2
        y = (self.winfo_screenheight() - 820)  // 2
        self.geometry(f"+{x}+{y}")
        self.grab_set()
        self.focus_force()

        # ── Shared state ──────────────────────────────────────────────────────
        self.made_var      = ctk.StringVar(value="")
        self.ans_vars:  dict[int, ctk.StringVar] = {}
        self.editor_status: ctk.CTkLabel | None  = None
        self.lbl_excel_status: ctk.CTkLabel | None = None

        # Tự luận: danh sách dòng động
        # Mỗi phần tử: {"tb": CTkTextbox(đáp án), "question_tb": CTkTextbox(câu hỏi), "diem_var": StringVar}
        self._tl_rows: list[dict] = []
        self._tl_scroll_inner: ctk.CTkFrame | None = None

        self._build_ui()

    # ═════════════════════════════════════════════════════════════════════════
    # UI SKELETON
    # ═════════════════════════════════════════════════════════════════════════

    def _build_ui(self):
        # ── Header ────────────────────────────────────────────────────────────
        hdr = ctk.CTkFrame(self, fg_color="transparent")
        hdr.pack(fill="x", padx=30, pady=(18, 4))
        ctk.CTkLabel(hdr, text="BỘ TẠO & QUẢN LÝ ĐÁP ÁN",
                     font=ctk.CTkFont(family="Roboto", size=24, weight="bold"),
                     text_color="#3498DB").pack(anchor="w")
        ctk.CTkLabel(hdr, text="Soạn thảo thủ công hoặc import từ file Excel hàng loạt",
                     font=ctk.CTkFont(size=13, slant="italic"),
                     text_color="gray60").pack(anchor="w", pady=(0,4))

        # ── Mã đề card (dùng chung 2 tab) ────────────────────────────────────
        card = ctk.CTkFrame(self, corner_radius=12, fg_color="gray20")
        card.pack(fill="x", padx=30, pady=(0, 6))
        inner = ctk.CTkFrame(card, fg_color="transparent")
        inner.pack(pady=12, padx=20, fill="x")

        ctk.CTkLabel(inner, text="MÃ ĐỀ:",
                     font=ctk.CTkFont(size=15, weight="bold")).pack(side="left", padx=(0,8))
        ctk.CTkEntry(inner, textvariable=self.made_var,
                     width=110, height=42, justify="center",
                     font=ctk.CTkFont(size=19, weight="bold"),
                     placeholder_text="001").pack(side="left", padx=6)
        ctk.CTkButton(inner, text="KIỂM TRA MÃ",
                      command=lambda: self._load_code(silent=False),
                      width=140, height=42,
                      font=ctk.CTkFont(size=13, weight="bold"),
                      fg_color="#34495E", hover_color="#2C3E50").pack(side="left", padx=14)
        self.editor_status = ctk.CTkLabel(
            inner, text="Nhập mã đề để tải hoặc tạo mới.",
            font=ctk.CTkFont(size=13, slant="italic"), text_color="gray60")
        self.editor_status.pack(side="left", padx=8)

        # ── Nút lưu (bottom — pack trước tabview) ─────────────────────────────
        bot = ctk.CTkFrame(self, fg_color="transparent")
        bot.pack(side="bottom", fill="x", pady=(8, 18))
        ctk.CTkButton(bot, text="💾  LƯU ĐÁP ÁN", command=self._save,
                      width=260, height=50, corner_radius=14,
                      fg_color="#2980B9", hover_color="#1F618D",
                      font=ctk.CTkFont(size=19, weight="bold")).pack()

        # ── TabView ───────────────────────────────────────────────────────────
        self.tabview = ctk.CTkTabview(self, corner_radius=12, fg_color="gray15",
                                      segmented_button_selected_color="#2980B9",
                                      segmented_button_selected_hover_color="#3498DB",
                                      segmented_button_unselected_color="gray25",
                                      segmented_button_unselected_hover_color="gray30")
        self.tabview.pack(fill="both", expand=True, padx=30, pady=(0, 6))

        self.tabview.add("📝  Trắc Nghiệm")
        self.tabview.add("✏️  Tự Luận")

        self._build_tn_tab(self.tabview.tab("📝  Trắc Nghiệm"))
        self._build_tl_tab(self.tabview.tab("✏️  Tự Luận"))

    # ═════════════════════════════════════════════════════════════════════════
    # TAB TRẮC NGHIỆM
    # ═════════════════════════════════════════════════════════════════════════

    def _build_tn_tab(self, tab):
        # Excel toolbar
        excel_bar = ctk.CTkFrame(tab, fg_color="gray22", corner_radius=10)
        excel_bar.pack(fill="x", padx=10, pady=(10, 6))
        ei = ctk.CTkFrame(excel_bar, fg_color="transparent")
        ei.pack(pady=10, padx=14, fill="x")
        ctk.CTkLabel(ei, text="📊  Excel:",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color="#2ecc71").pack(side="left", padx=(0,10))
        ctk.CTkButton(ei, text="⬇  Tải file mẫu",
                      command=self._download_template,
                      height=34, width=170, corner_radius=8,
                      font=ctk.CTkFont(size=12, weight="bold"),
                      fg_color="#27ae60", hover_color="#1e8449").pack(side="left", padx=(0,8))
        ctk.CTkButton(ei, text="📂  Import Excel",
                      command=self._import_from_excel,
                      height=34, width=160, corner_radius=8,
                      font=ctk.CTkFont(size=12, weight="bold"),
                      fg_color="#2980B9", hover_color="#1F618D").pack(side="left", padx=(0,8))
        self.lbl_excel_status = ctk.CTkLabel(
            ei, text="", font=ctk.CTkFont(size=11, slant="italic"), text_color="gray60")
        self.lbl_excel_status.pack(side="left", padx=6)

        # Grid 40 câu
        grid_frame = ctk.CTkFrame(tab, fg_color="gray18", corner_radius=10)
        grid_frame.pack(fill="both", expand=True, padx=10, pady=(0,10))
        for i in range(4):
            grid_frame.grid_columnconfigure(i, weight=1)
        self.after(50, lambda: self._build_tn_column(grid_frame, 0))

    def _build_tn_column(self, parent, col: int):
        if col >= 4:
            self._load_code(silent=True)
            return
        panel = ctk.CTkFrame(parent, fg_color="gray22", corner_radius=10)
        panel.grid(row=0, column=col, padx=8, pady=12, sticky="nsew")
        start_q = col * 10 + 1
        ctk.CTkLabel(panel, text=f"Câu {start_q} – {start_q+9}",
                     font=ctk.CTkFont(weight="bold", size=14),
                     text_color="#95A5A6").pack(pady=(10,8))
        for i in range(10):
            q_num = col * 10 + i + 1
            rf = ctk.CTkFrame(panel, fg_color="transparent")
            rf.pack(pady=3, fill="x", padx=12)
            ctk.CTkLabel(rf, text=f"{q_num:02d}", width=28,
                         font=ctk.CTkFont(weight="bold", size=13),
                         text_color="#BDC3C7").pack(side="left")
            var = ctk.StringVar(value="A")
            self.ans_vars[q_num] = var
            ctk.CTkSegmentedButton(rf, values=["A","B","C","D"], variable=var,
                                   selected_color="#2980B9", selected_hover_color="#3498DB",
                                   unselected_color="gray30",
                                   unselected_hover_color="gray35"
                                   ).pack(side="right", expand=True, fill="x")
        self.after(20, lambda: self._build_tn_column(parent, col+1))

    # ═════════════════════════════════════════════════════════════════════════
    # TAB TỰ LUẬN
    # ═════════════════════════════════════════════════════════════════════════

    def _build_tl_tab(self, tab):
        # Toolbar trên
        toolbar = ctk.CTkFrame(tab, fg_color="gray22", corner_radius=10)
        toolbar.pack(fill="x", padx=10, pady=(10, 6))
        ti = ctk.CTkFrame(toolbar, fg_color="transparent")
        ti.pack(pady=10, padx=14, fill="x")

        ctk.CTkLabel(ti, text="Số câu tự luận:",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color="gray80").pack(side="left", padx=(0,8))
        self.tl_count_var = ctk.StringVar(value=str(NUM_TL_DEFAULT))
        ctk.CTkEntry(ti, textvariable=self.tl_count_var,
                     width=60, height=34, justify="center",
                     font=ctk.CTkFont(size=14, weight="bold")).pack(side="left", padx=(0,8))
        ctk.CTkButton(ti, text="↺  Đặt lại số câu",
                      command=self._reset_tl_rows,
                      height=34, width=160, corner_radius=8,
                      font=ctk.CTkFont(size=12, weight="bold"),
                      fg_color="#7f8c8d", hover_color="#636e72").pack(side="left", padx=(0,12))

        self.lbl_tl_total = ctk.CTkLabel(
            ti, text="Tổng điểm TL: 0.0",
            font=ctk.CTkFont(size=13, weight="bold"), text_color="#e67e22")
        self.lbl_tl_total.pack(side="right", padx=8)

        # Header cột
        col_hdr = ctk.CTkFrame(tab, fg_color="gray25", corner_radius=8)
        col_hdr.pack(fill="x", padx=10, pady=(0,4))
        ctk.CTkLabel(col_hdr, text="Câu", width=50,
                     font=ctk.CTkFont(size=12, weight="bold"),
                     text_color="gray70").grid(row=0, column=0, padx=(14,4), pady=6)
        ctk.CTkLabel(col_hdr, text="Câu hỏi (ngữ cảnh cho AI)",
                     font=ctk.CTkFont(size=12, weight="bold"),
                     text_color="#e67e22").grid(row=0, column=1, padx=4, pady=6, sticky="w")
        ctk.CTkLabel(col_hdr, text="Đáp án mẫu",
                     font=ctk.CTkFont(size=12, weight="bold"),
                     text_color="gray70").grid(row=0, column=2, padx=4, pady=6, sticky="w")
        ctk.CTkLabel(col_hdr, text="Điểm", width=80,
                     font=ctk.CTkFont(size=12, weight="bold"),
                     text_color="gray70").grid(row=0, column=3, padx=(4,14), pady=6)
        col_hdr.grid_columnconfigure(1, weight=1)
        col_hdr.grid_columnconfigure(2, weight=1)

        # Scrollable area chứa các dòng câu hỏi
        self._tl_scroll = ctk.CTkScrollableFrame(tab, fg_color="gray18", corner_radius=10)
        self._tl_scroll.pack(fill="both", expand=True, padx=10, pady=(0,8))
        self._tl_scroll.grid_columnconfigure(1, weight=1)
        self._tl_scroll.grid_columnconfigure(2, weight=1)
        self._tl_scroll_inner = self._tl_scroll

        # Build dòng mặc định
        self.after(60, lambda: self._init_tl_rows(NUM_TL_DEFAULT))

    def _init_tl_rows(self, n: int):
        """Tạo n dòng trống cho tab TL."""
        # Xoá hết dòng cũ
        for widget in self._tl_scroll_inner.winfo_children():
            widget.destroy()
        self._tl_rows.clear()
        for i in range(n):
            self._add_tl_row(i + 1)

    def _add_tl_row(self, q_num: int):
        """Thêm 1 dòng câu tự luận vào scroll frame."""
        row = len(self._tl_rows)
        bg = "gray22" if row % 2 == 0 else "gray20"

        # Số câu
        ctk.CTkLabel(self._tl_scroll_inner, text=f"Câu {q_num}", width=50,
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color="#3498DB",
                     fg_color=bg, corner_radius=0
                     ).grid(row=row, column=0, padx=(10,4), pady=4, sticky="nsew")

        # Textbox câu hỏi (ngữ cảnh cho AI)
        q_tb = ctk.CTkTextbox(self._tl_scroll_inner, height=70, corner_radius=8,
                              font=ctk.CTkFont(size=12),
                              fg_color="gray30", text_color="#f0e68c",
                              border_color="#e67e22", border_width=1)
        q_tb.grid(row=row, column=1, padx=4, pady=4, sticky="ew")
        q_tb.insert("0.0", "")

        # Textbox đáp án mẫu
        tb = ctk.CTkTextbox(self._tl_scroll_inner, height=70, corner_radius=8,
                            font=ctk.CTkFont(size=12),
                            fg_color="gray28", text_color="white",
                            border_color="gray35", border_width=1)
        tb.grid(row=row, column=2, padx=4, pady=4, sticky="ew")
        tb.insert("0.0", "")

        # Entry điểm tối đa
        diem_var = ctk.StringVar(value="2.0")
        diem_entry = ctk.CTkEntry(
            self._tl_scroll_inner, textvariable=diem_var,
            width=80, height=36, justify="center",
            font=ctk.CTkFont(size=13, weight="bold"),
            placeholder_text="điểm")
        diem_entry.grid(row=row, column=3, padx=(4,10), pady=4)
        diem_var.trace_add("write", lambda *_: self._update_tl_total())

        self._tl_rows.append({"question_tb": q_tb, "tb": tb, "diem_var": diem_var})
        self._tl_scroll_inner.grid_columnconfigure(1, weight=1)
        self._tl_scroll_inner.grid_columnconfigure(2, weight=1)

    def _reset_tl_rows(self):
        """Đặt lại số câu theo ô nhập."""
        try:
            n = int(self.tl_count_var.get().strip())
            if n < 1 or n > 30:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Lỗi", "Số câu tự luận phải là số nguyên từ 1 đến 30.")
            return
        self._init_tl_rows(n)
        self._update_tl_total()

    def _update_tl_total(self):
        """Cập nhật nhãn tổng điểm TL."""
        total = 0.0
        for r in self._tl_rows:
            try:
                total += float(r["diem_var"].get())
            except ValueError:
                pass
        self.lbl_tl_total.configure(text=f"Tổng điểm TL: {total:.1f}")

    def _get_tl_list(self) -> list[dict]:
        """Đọc tất cả dòng TL thành list để lưu."""
        result = []
        for i, r in enumerate(self._tl_rows):
            cau_hoi = r["question_tb"].get("0.0", "end").strip()
            dap_an = r["tb"].get("0.0", "end").strip()
            try:
                diem = float(r["diem_var"].get())
            except ValueError:
                diem = 0.0
            result.append({"q": i + 1, "cau_hoi": cau_hoi, "dap_an": dap_an, "diem": diem})
        return result

    def _load_tl_rows(self, tl_list: list[dict]):
        """Điền dữ liệu từ db vào các dòng TL."""
        n = len(tl_list)
        if n != len(self._tl_rows):
            self._init_tl_rows(n)
            self.tl_count_var.set(str(n))
            # Đợi build xong rồi mới điền
            self.after(100, lambda: self._fill_tl_rows(tl_list))
        else:
            self._fill_tl_rows(tl_list)

    def _fill_tl_rows(self, tl_list: list[dict]):
        for i, item in enumerate(tl_list):
            if i >= len(self._tl_rows):
                break
            r = self._tl_rows[i]
            r["question_tb"].delete("0.0", "end")
            r["question_tb"].insert("0.0", item.get("cau_hoi", ""))
            r["tb"].delete("0.0", "end")
            r["tb"].insert("0.0", item.get("dap_an", ""))
            r["diem_var"].set(str(item.get("diem", 2.0)))
        self._update_tl_total()

    # ═════════════════════════════════════════════════════════════════════════
    # EXCEL ACTIONS (Trắc nghiệm)
    # ═════════════════════════════════════════════════════════════════════════

    def _download_template(self):
        path = filedialog.asksaveasfilename(
            title="Lưu file mẫu", defaultextension=".xlsx",
            filetypes=[("Excel files","*.xlsx")], initialfile="dap_an_mau.xlsx")
        if not path:
            return
        try:
            with open(path,"wb") as f:
                f.write(_create_template_bytes())
            self.lbl_excel_status.configure(
                text=f"✔ Đã lưu: {os.path.basename(path)}", text_color="#2ecc71")
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không thể lưu file mẫu:\n{e}")

    def _import_from_excel(self):
        path = filedialog.askopenfilename(
            title="Chọn file Excel đáp án",
            filetypes=[("Excel files","*.xlsx *.xls"),("All files","*.*")])
        if not path:
            return
        try:
            imported, errors = _import_excel(path)
        except ValueError as e:
            messagebox.showerror("Lỗi định dạng", str(e)); return
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không thể đọc file:\n{e}"); return
        if not imported:
            messagebox.showwarning("Không có dữ liệu","File không có hàng hợp lệ."); return

        summary = f"Tìm thấy {len(imported)} mã đề:\n{', '.join(sorted(imported.keys()))}"
        if errors:
            summary += f"\n\n⚠ {len(errors)} cảnh báo:\n" + "\n".join(errors[:8])
            if len(errors) > 8:
                summary += f"\n... và {len(errors)-8} cảnh báo khác"
        summary += "\n\nXác nhận import? (Mã đề trùng sẽ bị ghi đè)"
        if not messagebox.askyesno("Xác nhận Import", summary):
            return

        self.answer_key_db.update(imported)
        if self.on_bulk_save:
            self.on_bulk_save(self.answer_key_db)
        self.lbl_excel_status.configure(
            text=f"✔ Đã import {len(imported)} mã đề", text_color="#2ecc71")
        self._load_code(silent=True)
        messagebox.showinfo("Thành công", f"Đã import {len(imported)} mã đề!")

    # ═════════════════════════════════════════════════════════════════════════
    # LOGIC CHUNG
    # ═════════════════════════════════════════════════════════════════════════

    def _normalise_code(self) -> str | None:
        code = self.made_var.get().strip()
        if not code:
            return None
        if code.isdigit() and len(code) < 3:
            code = code.zfill(3)
            self.made_var.set(code)
        return code

    def _load_code(self, silent: bool = False):
        """Load đáp án TN + TL của mã đề hiện tại vào UI."""
        code = self._normalise_code()
        if not code:
            if not silent:
                messagebox.showwarning("Lỗi","Vui lòng nhập mã đề!")
            return

        found_tn = code in self.answer_key_db
        found_tl = code in self.tl_key_db

        # Điền TN
        if found_tn:
            for q, ans in self.answer_key_db[code].items():
                if q in self.ans_vars:
                    self.ans_vars[q].set(ans)

        # Điền TL
        if found_tl:
            self._load_tl_rows(self.tl_key_db[code])

        if found_tn or found_tl:
            parts = []
            if found_tn: parts.append("TN")
            if found_tl: parts.append("TL")
            self.editor_status.configure(
                text=f"Đã tải đáp án '{code}' ({' + '.join(parts)})",
                text_color="#3498DB")
        else:
            self.editor_status.configure(
                text=f"Mã '{code}' trống. Hãy tạo đáp án mới.",
                text_color="#95A5A6")

    def _save(self):
        code = self._normalise_code()
        if not code:
            messagebox.showwarning("Lỗi","Vui lòng nhập mã đề!")
            return

        tn_dict  = {q: self.ans_vars[q].get() for q in range(1, NUM_QUESTIONS+1)}
        tl_list  = self._get_tl_list()

        self.answer_key_db[code] = tn_dict
        self.tl_key_db[code]     = tl_list
        self.on_save(code, tn_dict, tl_list)
        messagebox.showinfo("Thành công", f"Đã lưu đáp án TN + TL cho mã đề '{code}'")